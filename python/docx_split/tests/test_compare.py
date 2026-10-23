"""compare 模块单测：归一化映射 / 句子拆分 / 端到端查重。"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from doc_analyze.compare import normalize as norm_mod
from doc_analyze.compare.engine import run_compare


# ---- normalize ----

def test_normalize_fullwidth_and_mapping_roundtrip():
    raw = "ＡＢＣ项目（重要）：提供稳定服务。"
    norm, mapping = norm_mod.normalize(raw)
    assert norm == "abc项目(重要):提供稳定服务."
    # 映射契约：len 一致、严格递增；拼回原始字符 = 原文（本例无删除）
    assert len(mapping) == len(norm)
    assert all(a < b for a, b in zip(mapping, mapping[1:]))
    assert "".join(raw[i] for i in mapping) == raw


def test_normalize_whitespace_collapse_and_strip():
    raw = "  提供\t稳定 \n 服务  "
    norm, mapping = norm_mod.normalize(raw)
    assert norm == "提供 稳定 服务"
    # 空白折叠：归一化空格映射到该连续空白段的第一个字符；首尾空白被抑制
    assert "".join(raw[i] for i in mapping) == "提供\t稳定 服务"


def test_normalize_zero_width_and_ctrl_dropped():
    raw = "提\u200b供\u0000服务"
    norm, mapping = norm_mod.normalize(raw)
    assert norm == "提供服务"
    assert len(mapping) == len(norm)
    assert "".join(raw[i] for i in mapping) == "提供服务"  # 零宽/控制字符被删除


# ---- split_sentences ----

def test_split_sentences_basic_and_digit_guard():
    norm, _ = norm_mod.normalize("提供稳定服务。版本号为2.1。交付验收标准严格执行！")
    spans = norm_mod.split_sentences(norm)
    texts = [norm[s:e] for s, e in spans]
    assert texts == ["提供稳定服务", "版本号为2.1", "交付验收标准严格执行"]


def test_split_sentences_filters_pure_numbering():
    norm, _ = norm_mod.normalize("一、项目概况。技术要求如下：")
    spans = norm_mod.split_sentences(norm)
    texts = [norm[s:e] for s, e in spans]
    assert "一、项目概况" not in texts
    assert "项目概况" in texts


# ---- 端到端 ----

def _make_doc(path: Path, title: str, paras: list[str], table_texts: list[list[str]] | None = None):
    d = Document()
    d.add_heading(title, level=1)
    for p in paras:
        d.add_paragraph(p)
    if table_texts:
        t = d.add_table(rows=len(table_texts), cols=len(table_texts[0]))
        for r, row in enumerate(table_texts):
            for c, txt in enumerate(row):
                t.cell(r, c).text = txt
    d.save(str(path))


@pytest.fixture(scope="module")
def docs_pair(tmp_path_factory):
    base = tmp_path_factory.mktemp("cmp")
    dup1 = "本项目采用成熟的微服务架构，系统性能稳定可靠，满足招标文件的全部技术要求。"
    dup2 = "投标人须在投标截止时间前，将投标文件密封送达指定开标地点，逾期不予受理。"
    dup_cell = "双电源冗余供电，自动切换时间小于10毫秒，保障系统持续稳定运行"
    uniq_a = "基准文档独有的技术方案描述：采用量子加密隧道保障数据传输安全。"
    uniq_b = "待查文档独有内容：售后响应时间为四小时以内，季度巡检不少于两次。"
    a = base / "a.docx"
    b = base / "b.docx"
    _make_doc(a, "基准标书", [dup1, uniq_a, dup2], [["序号", "参数"], ["1", dup_cell]])
    _make_doc(b, "待查标书", [dup1, uniq_b, dup2], [["序号", "参数"], ["1", dup_cell]])
    return a, b, dup1, dup2


def test_run_compare_finds_fragments(docs_pair):
    a, b, dup1, dup2 = docs_pair
    result = run_compare(a, b)
    frags = result["fragments"]
    assert frags, "至少应检出重复片段"

    contents = "".join(f["content"] for f in frags)
    assert "微服务架构" in contents
    assert "密封送达" in contents or "投标截止时间" in contents

    # 表格单元格重复（长单元格文本）应被检出且 match_type=表格
    table_frags = [f for f in frags if f["match_type"] == "表格"]
    assert table_frags and "双电源冗余" in "".join(f["content"] for f in table_frags)

    # 独有内容不应出现在片段中
    assert "量子加密" not in contents
    assert "售后响应" not in contents


def test_run_compare_offsets_align_with_block_text(docs_pair):
    a, b, *_ = docs_pair
    result = run_compare(a, b)

    # 从块树收集 block_id → 文本（段落直取；表格走 cells）
    def block_texts(tree):
        out = {}
        for node in tree:
            if node["kind"] == "paragraph":
                out[node["block"]] = node["text"]
            else:
                for cell in node["cells"]:
                    out[cell["block"]] = cell["text"]
        return out

    ta = block_texts(result["docs"]["a"]["tree"])
    tb = block_texts(result["docs"]["b"]["tree"])
    for f in result["fragments"]:
        la, lb = f["a"], f["b"]
        assert la["end"] > la["start"] and lb["end"] > lb["start"]
        # 块原始文本偏移切片 == content（同源构造性一致）
        assert ta[la["block"]][la["start"]:la["end"]] == f["content"]
        assert tb[lb["block"]][lb["start"]:lb["end"]] != "" or f["content"] == ""


def test_run_compare_stats(docs_pair):
    a, b, *_ = docs_pair
    result = run_compare(a, b)
    stats = result["stats"]
    assert stats["total_chars_a"] > 0 and stats["total_chars_b"] > 0
    assert 0 < stats["rate_b"] <= 1
    assert stats["fragment_count"] == len(result["fragments"])
    assert stats["by_confidence"]["high"]["fragments"] > 0
    assert stats["by_chapter"], "分章节统计非空"
    assert stats["elapsed_seconds"] >= 0


def test_run_compare_unique_docs_near_zero(tmp_path):
    a = tmp_path / "u1.docx"
    b = tmp_path / "u2.docx"
    _make_doc(a, "文档一", ["完全不同的第一份文档内容，讲述人工智能推理框架的设计与实现细节。"])
    _make_doc(b, "文档二", ["另一份毫无交集的文本，讨论海洋生物多样性与珊瑚礁保护策略研究。"])
    result = run_compare(a, b)
    assert result["stats"]["rate_b"] < 0.2
